# views.py
import csv
from pathlib import Path
from datetime import date, datetime
import re
from django.conf import settings
from django.http import JsonResponse, HttpResponse, HttpResponseBadRequest
from django.shortcuts import render, redirect, get_object_or_404
from django.core.cache import cache
from django.contrib import messages
from django.db import transaction
from django.db.models import Q, Sum, Min
from django.core.exceptions import ValidationError

from django_tables2 import RequestConfig
from openpyxl import Workbook

from collections import defaultdict
from decimal import Decimal
from itertools import chain

from project.decorators import menu_access_required, set_submenu_session
from core.helpers.filter_helper import FilterHelper

from sipd.registry import (
    SIPD_REGISTRY,
    get_rencana_by_pk,
    get_sp2d_sudah_realisasi_global
)

from .models import Sipd, TBP
from pendidikan.models import Realisasi
from .forms import SipdUploadForm, TBPUploadForm
from .tables import SipdTable, TBPTable
from .filters import SipdFilter
from .tasks import import_sipd_task, import_tbp_task


# =========================================================
# 🔧 GENERIC HELPERS
# =========================================================
def handle_uploaded_file(file, folder):
    upload_dir = Path(settings.MEDIA_ROOT) / folder
    upload_dir.mkdir(parents=True, exist_ok=True)

    file_path = upload_dir / file.name
    with open(file_path, "wb+") as f:
        for chunk in file.chunks():
            f.write(chunk)

    return file_path


def apply_pagination(request, table):
    per_page = request.GET.get("per_page", "10")
    if per_page != "all":
        RequestConfig(request, paginate={"per_page": int(per_page)}).configure(table)


def import_progress(request, task_id, prefix):
    data = cache.get(f"{prefix}_{task_id}") or {"current": 0, "total": 1}
    return JsonResponse(data)


def skipped_view(request, tahun, prefix):
    path = Path(settings.MEDIA_ROOT) / "import" / f"{prefix}_skipped_{tahun}.csv"

    rows = []
    if path.exists():
        with open(path, newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            rows = [{k.strip(): (v.strip() if v else "") for k, v in row.items()} for row in reader]

    return render(request, "components/skipped_list.html", {
        "tahun": tahun,
        "rows": rows,
    })


# =========================================================
# 📥 UPLOAD SIPD
# =========================================================
def upload_sipd(request):
    tahun = request.session.get("tahun")
    form = SipdUploadForm(request.POST or None, request.FILES or None)

    if request.method == "POST":
        if not form.is_valid():
            return JsonResponse({"error": form.errors}, status=400)

        file_path = handle_uploaded_file(form.cleaned_data["file"], "sipd")
        task = import_sipd_task.delay(str(file_path), tahun)

        return JsonResponse({"task_id": task.id})

    qs = Sipd.objects.all()
    if tahun:
        qs = qs.filter(tahun=tahun)

    sipd_filter = SipdFilter(request.GET, queryset=qs)
    table = SipdTable(sipd_filter.qs)
    apply_pagination(request, table)

    context = {
        "form": form,
        "table": table,
        "judul": "Upload Excel SIPD",
        "judulskip": "Data yang dilewati",
        "tahun": tahun,

        # 🔥 dynamic config
        "url_upload": "/sipd/",
        "url_progress": "/sipd/progress/TASK_ID/",
        "url_list": "/sipd/",
        "url_export": "/sipd/export/",
        "url_skipped": f"/sipd/skipped/{tahun}/" if tahun else None,

        "table_template": "components/partials/table.html",
    }

    template = "components/partials/table.html" if request.htmx else "components/upload.html"
    return render(request, template, context)


def sipd_import_progress(request, task_id):
    return import_progress(request, task_id, "sipd")


def skipped_sipd_view(request, tahun):
    return skipped_view(request, tahun, "sipd")


def export_sipd_excel(request):
    tahun = request.session.get("tahun")

    qs = Sipd.objects.all()
    if tahun:
        qs = qs.filter(tahun=tahun)

    qs = SipdFilter(request.GET, queryset=qs).qs

    wb = Workbook()
    ws = wb.active
    ws.title = "Data SIPD"

    headers = [
        "Tahun", "Nama Sub SKPD", "Kode Sub Kegiatan",
        "Kode Rekening", "Nomor SP2D", "Nilai Realisasi"
    ]
    ws.append(headers)

    for row in qs.values_list(
        "tahun", "nama_sub_skpd", "kode_sub_kegiatan",
        "kode_rekening", "nomor_sp2d", "nilai_realisasi"
    ):
        ws.append(row)

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="sipd_{tahun}.xlsx"'

    wb.save(response)
    return response


# =========================================================
# 📥 UPLOAD TBP
# =========================================================
def upload_tbp(request):
    tahun = request.session.get("tahun")
    form = TBPUploadForm(request.POST or None, request.FILES or None)

    if request.method == "POST":
        if not form.is_valid():
            return JsonResponse({"error": form.errors}, status=400)

        file_path = handle_uploaded_file(form.cleaned_data["file"], "tbp")
        task = import_tbp_task.delay(str(file_path), tahun)

        return JsonResponse({"task_id": task.id})

    qs = TBP.objects.all().order_by("-tanggal_tbp")
    table = TBPTable(qs)
    apply_pagination(request, table)

    context = {
        "form": form,
        "table": table,
        "judul": "Upload TBP",

        "url_upload": "/sipd/tbp/",
        "url_progress": "/sipd/tbp/progress/TASK_ID/",
        "url_list": "/sipd/tbp/",
        "url_export": "/sipd/tbp/export/",
        "url_skipped": None,

        "table_template": "components/partials/table.html",
    }

    template = "components/partials/table.html" if request.htmx else "components/upload.html"
    return render(request, template, context)


def tbp_import_progress(request, task_id):
    return import_progress(request, task_id, "tbp")


def export_tbp_excel(request):
    qs = TBP.objects.all().order_by("-tanggal_tbp")

    wb = Workbook()
    ws = wb.active
    ws.title = "Data TBP"

    ws.append(["Tahun", "Tanggal", "Nomor", "Nilai"])

    for row in qs.values_list("tahun", "tanggal_tbp", "nomor_tbp", "nilai_tbp"):
        ws.append(row)

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="tbp.xlsx"'

    wb.save(response)
    return response

# ================= HELPER =================
def clean(value):
    return str(value).strip() if value else ""


def extract_nomor_dari_status(status):
    """
    Ambil nomor setelah 'nomor:' dari teks seperti:
    - Telah Terbit SP2D nomor: xxx
    - Telah Terbit LPJ BP nomor: xxx
    """
    if not status:
        return ""

    match = re.search(r"nomor\s*:\s*([0-9./A-Z\-]+)", status, re.IGNORECASE)
    return clean(match.group(1)) if match else ""


def modal_output(request):
    key = request.POST.get('key') or request.GET.get('key')

    if key not in SIPD_REGISTRY:
        return HttpResponse("<div class='alert alert-danger'>Key tidak valid</div>")

    config = SIPD_REGISTRY[key]
    model_realisasi = config['model_realisasi']

    if request.method == "POST":
        try:
            pk = int(request.POST.get('id'))
        except (TypeError, ValueError):
            return HttpResponse("<div class='alert alert-danger'>ID tidak valid</div>")

        output = request.POST.get('realisasi_output')

        obj = get_object_or_404(model_realisasi, pk=pk)
        obj.realisasi_output = output

        try:
            obj.save()
        except Exception as e:
            return HttpResponse(f"<div class='alert alert-danger'>{e}</div>")

        return HttpResponse("""
            <script>
                location.reload();
            </script>
        """)

    # GET (load modal)
    pk = request.GET.get('id')
    obj = model_realisasi.objects.filter(pk=pk).first()

    return render(request, 'components/modal_output.html', {
        'obj': obj,
        'id': pk,
        'key': key,  # penting untuk dikirim balik ke POST
        
    })
    
    
# ================= VIEW =================
@set_submenu_session
@menu_access_required('list')
def view_sipd(request, mode, pk):
    request.session['next'] = request.get_full_path()

    config = SIPD_REGISTRY.get(mode)
    if not config:
        return HttpResponseBadRequest('Mode SIPD tidak valid')

    model_rencana = config['model_rencana']
    model_rencana_sisa = config['model_rencana_sisa']
    model_realisasi = config['model_realisasi']
    model_realisasi_sisa = config.get('model_realisasi_sisa')
    url_sp2d = config['url_sp2d']

    rencana = get_rencana_by_pk(pk, model_rencana, model_rencana_sisa)

    if not rencana:
        messages.error(request, 'Data rencana tidak ditemukan')
        return redirect('url_list_rencana')

    # ================= FILTER SIPD =================
    base_filter = Q(
        tahun=rencana.posting_tahun,
        kode_sub_kegiatan=rencana.posting_ket,
        kode_sub_skpd=rencana.posting_subopd.sub_opd_kode
    )

    sipd_qs_raw = Sipd.objects.filter(base_filter)

    # ================= TBP MAP =================
    tbp_map = {}

    tbp_docs = sipd_qs_raw.filter(
        jenis_dokumen="TBP"
    ).values_list("nomor_dokumen", flat=True)

    tbp_qs = TBP.objects.filter(
        tahun=rencana.posting_tahun,
        nomor_tbp__in=tbp_docs
    )

    for tbp in tbp_qs:
        nomor = extract_nomor_dari_status(tbp.status_tbp)

        if nomor:
            tbp_map[clean(tbp.nomor_tbp)] = nomor

    # ================= DATA SUDAH REALISASI =================
    sp2d_done = get_sp2d_sudah_realisasi_global(
        rencana,
        model_rencana,
        model_realisasi,
        model_rencana_sisa,
        model_realisasi_sisa,
    )

    # ================= AGREGASI SIPD =================
    sipd_qs = (
        sipd_qs_raw
        .values('nomor_sp2d', 'nomor_dokumen', 'jenis_dokumen')
        .annotate(
            tanggal_sp2d=Min('tanggal_sp2d'),
            nilai_realisasi=Sum('nilai_realisasi'),
        )
        .order_by('tanggal_sp2d', 'nomor_sp2d')
    )

    # ================= GROUP FINAL =================
    grouped_sp2d = {}

    for row in sipd_qs:

        # 🔥 tentukan nomor final
        if row["jenis_dokumen"] == "TBP":
            nomor_sp2d_final = tbp_map.get(clean(row["nomor_dokumen"]), "")
        else:
            nomor_sp2d_final = clean(row["nomor_sp2d"])

        if not nomor_sp2d_final:
            continue

        # skip jika sudah direalisasi
        if nomor_sp2d_final in sp2d_done:
            continue

        # ================= GROUP =================
        if nomor_sp2d_final not in grouped_sp2d:
            grouped_sp2d[nomor_sp2d_final] = {
                "nomor_sp2d_final": nomor_sp2d_final,
                "tanggal_sp2d": row["tanggal_sp2d"],
                "nilai_realisasi": Decimal(0),
            }

        grouped_sp2d[nomor_sp2d_final]["nilai_realisasi"] += (
            row["nilai_realisasi"] or Decimal(0)
        )

        # ambil tanggal paling awal
        if row["tanggal_sp2d"]:
            existing = grouped_sp2d[nomor_sp2d_final]["tanggal_sp2d"]
            if not existing or row["tanggal_sp2d"] < existing:
                grouped_sp2d[nomor_sp2d_final]["tanggal_sp2d"] = row["tanggal_sp2d"]

    # ================= FINAL DATA =================
    data_final = list(grouped_sp2d.values())

    # mapping untuk POST
    sp2d_mapping = {
        row["nomor_sp2d_final"]: row
        for row in data_final
    }

    # ================= POST =================
    if request.method == "POST":
        selected = request.POST.getlist('sp2d')

        if not selected:
            messages.warning(request, 'Tidak ada SP2D yang dipilih')
            return redirect(request.path)

        objs = []

        for sp2d in selected:
            row = sp2d_mapping.get(sp2d)

            if not row:
                continue

            objs.append(
                model_realisasi(
                    realisasi_dana_id=request.session.get('realisasi_dana'),
                    realisasi_tahap_id=request.session.get('realisasi_tahap'),
                    realisasi_subopd_id=request.session.get('realisasi_subopd'),

                    realisasi_rencanaposting=rencana,
                    realisasi_rencana=rencana.posting_rencanaid,
                    realisasi_subkegiatan=rencana.posting_subkegiatan,

                    realisasi_tahun=request.session.get('realisasi_tahun'),
                    realisasi_output=0,
                    realisasi_sp2d=sp2d,
                    realisasi_tgl=row['tanggal_sp2d'] or date.today(),
                    realisasi_nilai=row['nilai_realisasi'],
                )
            )

        try:
            with transaction.atomic():
                for obj in objs:
                    obj.save()
        except ValidationError as e:
            messages.error(request, f'Gagal menyimpan: {e.messages}')
            return redirect(url_sp2d, pk=rencana.pk)

        messages.success(request, f'{len(objs)} SP2D berhasil disimpan')
        return redirect(url_sp2d, pk=rencana.pk)

    # ================= RENDER =================
    return render(request, 'sipd/view_sipd.html', {
        'rencana': rencana,
        'data': data_final,
        'sp2d_sudah_realisasi': sp2d_done,
        
    })

def monitor_sipd(request):
    tahun = request.session.get('tahun') or datetime.now().year
    jadwal = request.session.get('jadwal')
    dana = request.GET.get('dana') or None

    # ================= AMBIL SEMUA RENCANA (OBJECT) =================
    all_rencana = []

    for key, config in SIPD_REGISTRY.items():
        model_rencana = config['model_rencana']

        filters = {
            'posting_tahun': tahun,
            'posting_jadwal_id': jadwal,
        }

        if dana:
            filters['posting_dana_id'] = dana

        qs = (
            model_rencana.objects
            .select_related(
                'posting_subopd',
                'posting_subkegiatan'
            )
            .filter(**filters)
        )

        for obj in qs:
            item = {
                'posting_tahun': obj.posting_tahun,
                'posting_ket': obj.posting_ket,
                'subopd_kode': obj.posting_subopd.sub_opd_kode if obj.posting_subopd else '',
                'subopd_nama': obj.get_subopd(),
                'subkegiatan_nama': obj.get_subkegiatan(),
                'posting_pagu': obj.posting_pagu or Decimal(0),
            }
            all_rencana.append(item)

    # ================= GROUPING =================
    grouped = defaultdict(lambda: {
        'posting_pagu': Decimal(0),
        'subopd_nama': '',
        'subkegiatan_nama': ''
    })

    for item in all_rencana:
        key = (
            item['posting_tahun'],
            clean(item['posting_ket']),
            clean(item['subopd_kode'])
        )

        grouped[key]['posting_pagu'] += item['posting_pagu']

        # isi sekali saja (hindari overwrite kosong)
        if not grouped[key]['subopd_nama'] and item['subopd_nama']:
            grouped[key]['subopd_nama'] = item['subopd_nama']

        if not grouped[key]['subkegiatan_nama'] and item['subkegiatan_nama']:
            grouped[key]['subkegiatan_nama'] = item['subkegiatan_nama']

    # ================= FINAL LIST =================
    data_rencana = []

    for key, val in grouped.items():
        tahun_k, ket, opd = key

        data_rencana.append({
            'posting_tahun': tahun_k,
            'posting_ket': ket,
            'posting_subopd_kode': opd,
            'subopd_nama': val['subopd_nama'],
            'subkegiatan_nama': val['subkegiatan_nama'],
            'posting_pagu': val['posting_pagu'],
        })

    # ================= SIPD (AGGREGATE) =================
    sipd_qs = (
        Sipd.objects
        .filter(tahun=tahun)
        .values(
            'tahun',
            'kode_sub_kegiatan',
            'kode_sub_skpd'
        )
        .annotate(
            nilai_realisasi=Sum('nilai_realisasi')  # pastikan field benar
        )
    )

    sipd_map = {
        (
            s['tahun'],
            clean(s['kode_sub_kegiatan']),
            clean(s['kode_sub_skpd'])
        ): s['nilai_realisasi']
        for s in sipd_qs
    }

    # ================= PROSES AKHIR =================
    hasil = []
    total_pagu = Decimal(0)
    total_sipd = Decimal(0)

    for item in data_rencana:
        key = (
            item['posting_tahun'],
            clean(item['posting_ket']),
            clean(item['posting_subopd_kode'])
        )

        nilai_sipd = sipd_map.get(key, Decimal(0))
        selisih = item['posting_pagu'] - nilai_sipd

        hasil.append({
            **item,
            'nilai_sipd': nilai_sipd,
            'selisih': selisih,
            'status': 'SESUAI' if selisih == 0 else ('BELUM ADA' if nilai_sipd == 0 else 'SELISIH')
        })

        total_pagu += item['posting_pagu']
        total_sipd += nilai_sipd

    total_selisih = total_pagu - total_sipd

    # ================= CONTEXT =================
    context = {
        'data': hasil,
        'total_pagu': total_pagu,
        'total_sipd': total_sipd,
        'total_selisih': total_selisih,
        'tahun': tahun,
        'dana': dana,
        'dana_choices': FilterHelper.dana_choices(),
    }

    # ================= HTMX =================
    if request.headers.get('HX-Request'):
        return render(request, 'components/partials/table_monitor.html', context)

    return render(request, 'sipd/monitor_sipd.html', context)